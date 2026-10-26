"""Готовые кворки клиента: папка на Яндекс Диске (папка с фото + таблица Excel/CSV) → кворки.

Шаги (бот вызывает их по очереди, между ними — вопросы клиенту):
1. prepare()        — скачать, разобрать таблицу, сопоставить фото, проверить правила и уникальность;
2. apply_defaults() — цена/срок для строк, где их нет (клиент выбирает кнопкой);
3. finish()         — рубрики (ИИ), обложки из фото (при выборе «уникализировать» — обработка
                      похожих), сохранить кворки.
Состояние между шагами — в profile["bulk"].
"""

from __future__ import annotations

import csv
import io
import random
import re
import shutil
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from kwork_agent.config import CONFIG
from kwork_agent.cover import SIZE
from kwork_agent.generator import validate
from kwork_agent.uniqueness import find_duplicate

from . import content, db, disk, images
from .kwork_rules import category_list, find_category, fit_instruction, normalize

PHOTO_EXT = {".jpg", ".jpeg", ".png", ".webp"}
TABLE_EXT = {".xlsx", ".csv"}
PHOTO_DIRS = ("фото", "photo", "photos", "картинки", "изображения", "images", "обложки", "covers")
TEXT_DIRS = ("тексты", "текст", "texts", "text", "таблица")

# Какие слова в заголовке столбца что значат (сравниваем по началу/вхождению, без регистра).
COLUMNS = [
    ("photo", ("фото", "картин", "изображ", "обложк", "photo", "image")),
    ("requirements", ("что нужно", "покупател", "требован", "для начала")),
    ("title", ("назв", "заголов", "title")),
    ("description", ("опис", "текст", "description")),
    ("price_rub", ("цен", "стоим", "price")),
    ("days", ("срок", "дн", "days")),
    ("category_hint", ("рубрик", "категор", "раздел", "category")),
    ("tags", ("тег", "tag", "ключев")),
    ("volume", ("объём", "объем", "volume")),
]

CATEGORY_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["items"],
    "properties": {"items": {"type": "array", "items": {
        "type": "object", "additionalProperties": False, "required": ["row", "category_hint", "volume", "attribute_hints"],
        "properties": {"row": {"type": "integer"},
                       "category_hint": {"type": "string", "description": "Рубрика Kwork «Раздел > Подраздел»"},
                       "volume": {"type": "string", "description": "Объём услуг в 1 кворке: что входит в цену, "
                                                                   "с числом, до 60 символов"},
                       "attribute_hints": {"type": "array", "items": {"type": "string"},
                                           "description": "3–8 слов для полей рубрики: платформа, язык, программа"}}}}},
}


def work_dir(uid: int) -> Path:
    return db.user_dir(uid) / "bulk"


# --- Таблица ---------------------------------------------------------------------------------

def _column_map(header: list[str]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for i, raw in enumerate(header):
        name = str(raw or "").strip().lower()
        for key, words in COLUMNS:
            if key not in mapping.values() and any(w in name for w in words):
                mapping[i] = key
                break
    return mapping


def _int(value) -> int | None:
    digits = re.sub(r"[^\d]", "", str(value or "").split(",")[0].split(".")[0])
    return int(digits) if digits else None


def read_table(path: Path) -> list[dict]:
    """Строки таблицы → [{"row", "title", "description", ...}]. Пустые строки пропускаем."""
    if path.suffix.lower() == ".xlsx":
        import openpyxl

        sheet = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    else:
        raw = path.read_bytes()
        text = next((raw.decode(enc) for enc in ("utf-8-sig", "cp1251") if _decodes(raw, enc)), raw.decode("latin-1"))
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t") if text.strip() else csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        raise ValueError("таблица пустая")
    mapping = _column_map(rows[0])
    if "title" not in mapping.values() or "description" not in mapping.values():
        raise ValueError("в первой строке таблицы нужны столбцы «Название» и «Описание» — скачайте шаблон")
    out = []
    for n, row in enumerate(rows[1:], start=2):
        item = {key: row[i] for i, key in mapping.items() if i < len(row) and row[i] not in (None, "")}
        if not item.get("title") and not item.get("description"):
            continue
        out.append({
            "row": n,
            "title": str(item.get("title", "")).strip(),
            "description": str(item.get("description", "")).strip(),
            "requirements": str(item.get("requirements", "")).strip(),
            "price_rub": _int(item.get("price_rub")),
            "days": _int(item.get("days")),
            "category_hint": str(item.get("category_hint", "")).strip(),
            "tags": [t.strip() for t in re.split(r"[,;#]", str(item.get("tags", ""))) if t.strip()],
            "photo": str(item.get("photo", "")).strip(),
            "volume": str(item.get("volume", "")).strip(),
        })
    return out


def _decodes(raw: bytes, enc: str) -> bool:
    try:
        raw.decode(enc)
        return True
    except UnicodeDecodeError:
        return False


def template_xlsx() -> bytes:
    """Шаблон таблицы для клиента."""
    import openpyxl
    from openpyxl.styles import Alignment, Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Кворки"
    ws.append(["Фото", "Название", "Описание", "Что нужно от покупателя", "Цена", "Срок, дней", "Рубрика", "Теги",
               "Объём услуг"])
    ws.append(["1.jpg", "Разработаю логотип для вашего бизнеса",
               "Создам уникальный логотип с учётом вашей сферы и аудитории.\n\nВы получите:\n— 3 варианта\n— исходники",
               "Название компании, сфера, пожелания по цветам", 3000, 3, "Дизайн > Логотипы", "логотип, дизайн",
               "1 логотип в 3 вариантах"])
    ws.append(["2.jpg", "Сделаю дизайн визитки", "…", "", "", "", "", "", ""])
    for col, width in zip("ABCDEFGHI", (10, 40, 60, 35, 10, 12, 25, 25, 25)):
        ws.column_dimensions[col].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Скачивание и сопоставление -------------------------------------------------------------

def _find(items: list[dict], names: tuple[str, ...]) -> dict | None:
    return next((i for i in items if i["type"] == "dir" and i["name"].strip().lower() in names), None)


def _is_photo(item: dict) -> bool:
    return item["type"] == "file" and Path(item["name"]).suffix.lower() in PHOTO_EXT


def _is_table(item: dict) -> bool:
    return item["type"] == "file" and Path(item["name"]).suffix.lower() in TABLE_EXT


def match_photos(rows: list[dict], photo_names: list[str]) -> list[str]:
    """Сопоставляет строкам фото. Есть столбец «Фото» — по имени файла (можно без расширения),
    иначе по номеру (строка 1 ↔ 1.jpg), иначе по порядку. Возвращает список проблем."""
    by_name = {n.lower(): n for n in photo_names}
    by_stem = {Path(n).stem.lower(): n for n in photo_names}
    ordered = sorted(photo_names, key=lambda n: (_int(Path(n).stem) is None, _int(Path(n).stem) or 0, n.lower()))
    problems = []
    use_column = any(r["photo"] for r in rows)
    for k, r in enumerate(rows, start=1):
        if use_column:
            want = r["photo"].lower()
            r["photo"] = by_name.get(want) or by_stem.get(Path(want).stem) or ""
        else:
            r["photo"] = by_stem.get(str(k)) or (ordered[k - 1] if k - 1 < len(ordered) else "")
        if not r["photo"]:
            problems.append(f"строка {r['row']}: не нашёл фото")
    return problems


def prepare(uid: int, link: str, progress: dict) -> dict:
    """Скачивает папку, разбирает таблицу, проверяет. Результат кладёт в profile["bulk"]."""
    progress.update(stage="bulk_list", done=0, total=0)
    root = disk.list_dir(link)
    photos_dir = _find(root, PHOTO_DIRS)
    photo_items = disk.list_dir(link, photos_dir["path"]) if photos_dir else [i for i in root if _is_photo(i)]
    photo_items = [i for i in photo_items if _is_photo(i)]
    texts_dir = _find(root, TEXT_DIRS)
    tables = [i for i in root if _is_table(i)] + ([i for i in disk.list_dir(link, texts_dir["path"]) if _is_table(i)]
                                                    if texts_dir else [])
    if not tables:
        raise ValueError("не нашёл таблицу (.xlsx или .csv) — положите её в папку рядом с папкой «фото»")
    if not photo_items:
        raise ValueError("не нашёл фото — положите их в папку «фото» (jpg, png или webp)")

    wd = work_dir(uid)
    shutil.rmtree(wd, ignore_errors=True)
    table = disk.download(tables[0], wd / f"table{Path(tables[0]['name']).suffix.lower()}")
    rows = read_table(table)
    problems = match_photos(rows, [i["name"] for i in photo_items])

    needed = {r["photo"] for r in rows if r["photo"]}
    progress.update(stage="bulk_photos", done=0, total=len(needed))
    files = {i["name"]: i for i in photo_items}

    def fetch(name: str) -> None:
        raw = disk.download(files[name], wd / "raw" / name)
        _fit(raw, wd / "photos" / f"{Path(name).stem}.jpg")
        progress["done"] += 1

    names = sorted(needed)
    with ThreadPoolExecutor(max_workers=8) as pool:
        for name, result in zip(names, pool.map(lambda n: _safe(fetch, n), names)):
            if result:
                problems.append(f"фото «{name}»: {result}")
                for r in rows:
                    if r["photo"] == name:
                        r["photo"] = ""
    shutil.rmtree(wd / "raw", ignore_errors=True)

    rows, more = check_rows(uid, [r for r in rows if r["photo"]])
    problems += more
    free = content.room(uid)
    over = max(0, len(rows) - free)
    rows = rows[:free]
    state = {"rows": rows, "problems": problems, "over_limit": over,
             "twins": photo_twins(uid, rows), "photo_choice": None}
    profile = db.load_profile(uid)
    profile["bulk"] = state
    db.save_profile(uid, profile)
    return state


def _safe(fn, *args) -> str:
    try:
        fn(*args)
        return ""
    except Exception as e:  # noqa: BLE001
        return str(e)[:120] if isinstance(e, ValueError) else "не скачалось или не открылось"


def _fit(raw: Path, out: Path) -> Path:
    """Фото клиента → JPEG с запасом для обработки (обложка Kwork 660×440, соотношение 3:2)."""
    from PIL import Image, ImageOps

    out.parent.mkdir(parents=True, exist_ok=True)
    ImageOps.fit(Image.open(raw).convert("RGB"), (SIZE[0] * 2, SIZE[1] * 2), Image.LANCZOS).save(out, "JPEG", quality=92)
    return out


def photo_path(uid: int, row: dict) -> Path:
    return work_dir(uid) / "photos" / f"{Path(row['photo']).stem}.jpg"


# --- Проверки -----------------------------------------------------------------------------------

def check_rows(uid: int, rows: list[dict]) -> tuple[list[dict], list[str]]:
    """Правила Kwork и уникальность текстов (между строками и с прошлыми кворками клиента).
    Цена и срок могут быть пустыми — их спросим кнопкой."""
    problems, ok = [], []
    previous = content.previous_texts(uid)
    for r in rows:
        probe = {**r, "price_rub": r["price_rub"] or CONFIG.price_min, "days": r["days"] or 1}
        probe["requirements"] = fit_instruction(r["requirements"])  # короткое дополним сами
        errors = validate(probe)
        if r["price_rub"] and not CONFIG.price_min <= r["price_rub"] <= CONFIG.price_max:
            errors.append(f"цена вне {CONFIG.price_min}–{CONFIG.price_max} ₽")
        dup = find_duplicate(r, previous + ok)
        if dup:
            errors.append(f"текст похож на «{dup['title'][:50]}»")
        if errors:
            problems.append(f"строка {r['row']} «{r['title'][:40]}»: {'; '.join(errors)}")
        else:
            ok.append(r)
    return ok, problems


def photo_twins(uid: int, rows: list[dict]) -> list[list[int]]:
    """Пары похожих фото (номера строк) — между собой и с обложками прошлых партий клиента."""
    marks = {r["row"]: images.cover_fingerprint(photo_path(uid, r)) for r in rows}
    old = [h["fp"] for h in db.load_history(uid) if h.get("fp") is not None]
    twins, seen = [], set()
    rows_sorted = sorted(marks)
    for i, a in enumerate(rows_sorted):
        for b in rows_sorted[i + 1:]:
            if images.covers_similar(marks[a], marks[b]):
                twins.append([a, b])
                seen.add(b)
        if a not in seen and any(images.covers_similar(marks[a], o) for o in old):
            twins.append([a, 0])  # 0 — похоже на обложку из прошлой партии
    return twins


def missing(state: dict) -> dict[str, int]:
    rows = state["rows"]
    return {"price": sum(1 for r in rows if not r["price_rub"]), "days": sum(1 for r in rows if not r["days"])}


def apply_defaults(uid: int, price: int | None = None, days: int | None = None) -> dict:
    profile = db.load_profile(uid)
    state = profile["bulk"]
    for r in state["rows"]:
        if price and not r["price_rub"]:
            r["price_rub"] = price
        if days and not r["days"]:
            r["days"] = days
    db.save_profile(uid, profile)
    return state


# --- Финал ------------------------------------------------------------------------------------

def uniquify_photo(src: Path, seed: str, strength: float):
    """Мягкая обработка фото клиента: кадр, лёгкий поворот, оттенок, яркость, контраст, зерно.
    strength растёт с каждой попыткой, если фото всё ещё похоже на другое."""
    from PIL import Image, ImageEnhance

    r = random.Random(seed)
    w, h = SIZE
    img = Image.open(src).convert("RGB")
    zoom = 1 + r.uniform(0.04, 0.12) * strength
    cw, ch = int(img.width / zoom), int(img.height / zoom)
    x, y = r.randint(0, img.width - cw), r.randint(0, img.height - ch)
    img = img.crop((x, y, x + cw, y + ch)).rotate(r.uniform(-2, 2) * strength, resample=Image.BICUBIC)
    m = int(min(img.size) * 0.03 * strength)  # обрезаем углы после поворота
    img = img.crop((m, m, img.width - m, img.height - m)).resize((w, h), Image.LANCZOS)
    hue, sat, val = img.convert("HSV").split()
    shift = int(r.uniform(-12, 12) * strength)
    img = Image.merge("HSV", (hue.point(lambda p: (p + shift) % 256), sat, val)).convert("RGB")
    img = ImageEnhance.Brightness(img).enhance(1 + r.uniform(-0.08, 0.08) * strength)
    img = ImageEnhance.Contrast(img).enhance(1 + r.uniform(-0.08, 0.1) * strength)
    return Image.blend(img, Image.effect_noise((w, h), 12).convert("RGB"), min(0.06, 0.02 * strength))


def classify(rows: list[dict], llm=None) -> None:
    """Подбирает рубрику Kwork и «объём услуг в 1 кворке» строкам, где их нет (одним-двумя запросами к ИИ)."""
    todo = [r for r in rows if not find_category(r["category_hint"]) or not r.get("volume") or not r.get("attribute_hints")]
    if not todo:
        return
    llm = llm or content._llm()
    for i in range(0, len(todo), 40):
        chunk = todo[i:i + 40]
        listing = "\n".join(f"{r['row']}. {r['title']} — {r['description'][:150]}" for r in chunk)
        data = llm.ask("Подбери каждому кворку рубрику СТРОГО из списка рубрик Kwork (строку как есть) и объём "
                       "услуг в 1 кворке — что входит в цену, коротко и с числом («1 логотип в 3 вариантах», "
                       "«до 1000 символов текста»), и 3–8 слов для полей рубрики Kwork: платформа, язык "
                       "программирования, программа (Telegram, Python, Excel) — только то, что есть в услуге.\n\n"
                       f"КВОРКИ:\n{listing}\n\nРУБРИКИ KWORK:\n{category_list()}",
                       system=content.SYSTEM, schema=CATEGORY_SCHEMA, effort="low", max_tokens=8000)
        found = {it["row"]: it for it in data["items"]}
        for r in chunk:
            got = found.get(r["row"], {})
            if not find_category(r["category_hint"]):
                r["category_hint"] = got.get("category_hint", "")
            if not r.get("volume"):
                r["volume"] = str(got.get("volume", "")).strip()
            if not r.get("attribute_hints"):
                r["attribute_hints"] = [str(h) for h in got.get("attribute_hints") or []]


def own_photo_path(uid: int, lid: str) -> Path:
    return db.user_dir(uid) / "photos" / f"{lid}.jpg"


def make_unique(src: Path, out: Path, lid: str, taken: list[int], start: int) -> int | None:
    """Обрабатывает фото, пока его обложка похожа на какую-то из taken. Возвращает отпечаток
    (или None, если фото и так уникально и не трогалось)."""
    mark = images.cover_fingerprint(out)
    if not any(images.covers_similar(mark, o) for o in taken) and start == 1:
        return None
    for attempt in range(start, start + 15):
        images.save_cover_jpeg(uniquify_photo(src, f"{lid}:{attempt}", 1 + attempt / 3), out)
        mark = images.cover_fingerprint(out)
        if not any(images.covers_similar(mark, o) for o in taken):
            break
    return mark


def new_look_photo(uid: int, listing: dict, listings: list[dict]) -> None:
    """«Другой фон» у кворка с фото клиента: другая мягкая обработка его фото."""
    taken = [h["fp"] for h in db.load_history(uid) if h.get("fp") is not None]
    taken += [o["fp"] for o in listings if o["id"] != listing["id"] and "fp" in o]
    listing["variant"] = listing.get("variant", 0) + 1
    out = db.cover_path(uid, listing["id"])
    images.save_cover_jpeg(uniquify_photo(own_photo_path(uid, listing["id"]), f"{listing['id']}:v{listing['variant']}", 1.5), out)
    listing["fp"] = make_unique(own_photo_path(uid, listing["id"]), out, listing["id"], taken,
                                start=listing["variant"] * 20) or images.cover_fingerprint(out)


def finish(uid: int, progress: dict, llm=None) -> int:
    """Собирает кворки: рубрики, обложки (уникализация похожих — если клиент её выбрал)."""
    from PIL import Image, ImageOps

    profile = db.load_profile(uid)
    state = profile["bulk"]
    rows = state["rows"]
    try:
        classify(rows, llm)
    except Exception as e:  # noqa: BLE001 — без рубрики кворк всё равно можно выложить
        progress["error"] = f"рубрики: {type(e).__name__}: {e}"[:300]
    content.reset(uid)  # прошлая партия → история; дальше — новая
    profile = db.load_profile(uid)
    fix = state.get("photo_choice") == "uniquify"
    taken = [h["fp"] for h in db.load_history(uid) if h.get("fp") is not None]
    listings = []
    progress.update(stage="covers", done=0, total=len(rows))
    for n, r in enumerate(rows, start=1):
        lid = uuid.uuid4().hex[:10]
        src = own_photo_path(uid, lid)  # оригинал храним: «Другой фон» обработает его заново
        src.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(photo_path(uid, r), src)
        out = db.cover_path(uid, lid)
        images.save_cover_jpeg(ImageOps.fit(Image.open(src).convert("RGB"), SIZE, Image.LANCZOS), out)
        mark = images.cover_fingerprint(out)
        if fix:
            mark = make_unique(src, out, lid, taken, start=1) or mark
        taken.append(mark)
        listing = {
            "id": lid, "idea_no": n, "idea": r["title"], "title": r["title"], "description": r["description"],
            "requirements": r["requirements"], "price_rub": r["price_rub"],
            "days": r["days"], "category_hint": r["category_hint"], "tags": r["tags"],
            "volume": r.get("volume", ""), "attribute_hints": r.get("attribute_hints", []),
            "cover_text": r["title"], "own_photo": True, "fp": mark, "row": r["row"],
        }
        normalize(listing)  # рубрика → номера Kwork, цена из допустимых, «что нужно» ≥ 100 символов
        listings.append(listing)
        progress["done"] = n
    db.save_listings(uid, listings)
    profile = db.load_profile(uid)
    profile.pop("bulk", None)
    profile["ideas"] = []
    db.save_profile(uid, profile)
    shutil.rmtree(work_dir(uid), ignore_errors=True)
    return len(listings)

